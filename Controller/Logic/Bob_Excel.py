######################################################
# Bob Excel Service
#
#
######################################################

# Python Imports

import logging
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import List, Dict, Any, Optional, Tuple


######################## Bob Excel Service #####################

class Bob_Excel_Service:

    """
    Bob Excel Service - Creates and manipulates xlsx files
    """

    ########################## Inits #######################

    def __init__(self):
        """
        Inits
        """
        logging.debug('Initializing Bob Excel Service ..')
        self._workbook: Optional[Workbook] = None
        self._current_sheet = None

    ########################## Create New Workbook #######################

    def create_workbook(self):
        """
        Create a new workbook
        """
        self._workbook = Workbook()
        # Remove default sheet
        if 'Sheet' in self._workbook.sheetnames:
            del self._workbook['Sheet']
        return self

    ########################## Create New Sheet #######################

    def create_sheet(self, sheet_name: str):
        """
        Create a new sheet in the workbook
        :param sheet_name: Name of the sheet
        """
        if self._workbook is None:
            self.create_workbook()
        self._current_sheet = self._workbook.create_sheet(title=sheet_name)
        return self

    ########################## Get Sheet #######################

    def get_sheet(self, sheet_name: str):
        """
        Get an existing sheet by name
        :param sheet_name: Name of the sheet
        """
        if self._workbook is None:
            raise ValueError("No workbook created. Call create_workbook() first.")
        self._current_sheet = self._workbook[sheet_name]
        return self

    ########################## Add Table Headers #######################

    def add_table_headers(self, headers: List[str], row: int = 1,
                          bg_color: str = "4472C4", font_color: str = "FFFFFF",
                          bold: bool = True):
        """
        Add table headers with styling
        :param headers: List of header names
        :param row: Row number to add headers (1-indexed)
        :param bg_color: Background color hex (without #)
        :param font_color: Font color hex (without #)
        :param bold: Whether to bold the headers
        """
        if self._current_sheet is None:
            raise ValueError("No sheet selected. Call create_sheet() or get_sheet() first.")

        fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
        font = Font(bold=bold, color=font_color)
        alignment = Alignment(horizontal="center", vertical="center")
        border = Side(style='thin', color='000000')
        cell_border = Border(left=border, right=border, top=border, bottom=border)

        for col_idx, header in enumerate(headers, start=1):
            cell = self._current_sheet.cell(row=row, column=col_idx, value=header)
            cell.fill = fill
            cell.font = font
            cell.alignment = alignment
            cell.border = cell_border

        return self

    ########################## Add Row #######################

    def add_row(self, values: List[Any], row: int,
                bg_color: Optional[str] = None,
                font_color: str = "000000",
                bold: bool = False,
                align_center: bool = False):
        """
        Add a row of values
        :param values: List of values for each cell
        :param row: Row number (1-indexed)
        :param bg_color: Optional background color hex (without #)
        :param font_color: Font color hex (without #)
        :param bold: Whether to bold the text
        :param align_center: Whether to center align
        """
        if self._current_sheet is None:
            raise ValueError("No sheet selected. Call create_sheet() or get_sheet() first.")

        font = Font(bold=bold, color=font_color)
        border = Side(style='thin', color='000000')
        cell_border = Border(left=border, right=border, top=border, bottom=border)

        for col_idx, value in enumerate(values, start=1):
            cell = self._current_sheet.cell(row=row, column=col_idx, value=value)
            cell.font = font
            cell.border = cell_border
            if bg_color:
                cell.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
            if align_center:
                cell.alignment = Alignment(horizontal="center", vertical="center")

        return self

    ########################## Add Multiple Rows #######################

    def add_rows(self, data: List[List[Any]], start_row: int = 2):
        """
        Add multiple rows of data
        :param data: List of rows, each row is a list of values
        :param start_row: Starting row number (1-indexed)
        """
        for row_idx, row_data in enumerate(data, start=start_row):
            self.add_row(row_data, row=row_idx)
        return self

    ########################## Add Table #######################

    def add_table(self, headers: List[str], data: List[List[Any]],
                  start_row: int = 1,
                  header_bg_color: str = "4472C4",
                  header_font_color: str = "FFFFFF"):
        """
        Add a complete table with headers and data
        :param headers: List of header names
        :param data: List of rows
        :param start_row: Starting row for headers (1-indexed)
        :param header_bg_color: Header background color hex
        :param header_font_color: Header font color hex
        """
        self.add_table_headers(headers, row=start_row,
                               bg_color=header_bg_color,
                               font_color=header_font_color)
        self.add_rows(data, start_row=start_row + 1)
        return self

    ########################## Set Column Width #######################

    def set_column_width(self, column: int, width: float):
        """
        Set width for a specific column
        :param column: Column number (1-indexed)
        :param width: Width value
        """
        if self._current_sheet is None:
            raise ValueError("No sheet selected.")
        col_letter = get_column_letter(column)
        self._current_sheet.column_dimensions[col_letter].width = width
        return self

    ########################## Auto Fit Columns #######################

    def auto_fit_columns(self, min_width: float = 10, max_width: float = 50):
        """
        Auto-fit column widths based on content
        :param min_width: Minimum column width
        :param max_width: Maximum column width
        """
        if self._current_sheet is None:
            raise ValueError("No sheet selected.")

        for column_cells in self._current_sheet.columns:
            max_length = 0
            column_letter = get_column_letter(column_cells[0].column)
            for cell in column_cells:
                try:
                    cell_length = len(str(cell.value)) if cell.value else 0
                    if cell_length > max_length:
                        max_length = cell_length
                except:
                    pass
            adjusted_width = min(max(max_length + 2, min_width), max_width)
            self._current_sheet.column_dimensions[column_letter].width = adjusted_width

        return self

    ########################## Auto Fit Rows #######################

    def auto_fit_rows(self, min_height: float = 15, max_height: float = 100,
                      default_col_width: float = 10):
        """
        Auto-fit row heights based on content (handles text wrapping)
        :param min_height: Minimum row height
        :param max_height: Maximum row height
        :param default_col_width: Default column width for calculating wrap lines
        """
        if self._current_sheet is None:
            raise ValueError("No sheet selected.")

        for row_idx, row_cells in enumerate(self._current_sheet.iter_rows(), start=1):
            max_lines = 1
            for cell in row_cells:
                if cell.value:
                    cell_text = str(cell.value)
                    # Get column width
                    col_letter = get_column_letter(cell.column)
                    col_width = self._current_sheet.column_dimensions[col_letter].width or default_col_width
                    # Estimate lines needed (rough: ~1 char per unit width)
                    chars_per_line = max(int(col_width), 1)
                    # Count explicit newlines + wrapped lines
                    lines = 1
                    for line in cell_text.split('\n'):
                        lines += max(0, (len(line) - 1) // chars_per_line)
                    lines += cell_text.count('\n')
                    if lines > max_lines:
                        max_lines = lines
            # ~15 points per line is typical
            calculated_height = max_lines * 15
            adjusted_height = min(max(calculated_height, min_height), max_height)
            self._current_sheet.row_dimensions[row_idx].height = adjusted_height

        return self

    ########################## Set Row Height #######################

    def set_row_height(self, row: int, height: float):
        """
        Set height for a specific row
        :param row: Row number (1-indexed)
        :param height: Height value in points
        """
        if self._current_sheet is None:
            raise ValueError("No sheet selected.")
        self._current_sheet.row_dimensions[row].height = height
        return self

    ########################## Set Cell Value #######################

    def set_cell(self, row: int, column: int, value: Any,
                 bg_color: Optional[str] = None,
                 font_color: str = "000000",
                 bold: bool = False):
        """
        Set a single cell value with optional styling
        :param row: Row number (1-indexed)
        :param column: Column number (1-indexed)
        :param value: Cell value
        :param bg_color: Optional background color hex
        :param font_color: Font color hex
        :param bold: Whether to bold
        """
        if self._current_sheet is None:
            raise ValueError("No sheet selected.")

        cell = self._current_sheet.cell(row=row, column=column, value=value)
        cell.font = Font(bold=bold, color=font_color)
        if bg_color:
            cell.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")

        return self

    ########################## Save Workbook #######################

    def save(self, file_path: str):
        """
        Save the workbook to a file
        :param file_path: Path to save the xlsx file
        """
        if self._workbook is None:
            raise ValueError("No workbook to save. Call create_workbook() first.")
        self._workbook.save(file_path)
        logging.info('Saved workbook to: {}'.format(file_path))
        return self

    ########################## Get Workbook #######################

    def get_workbook(self) -> Workbook:
        """
        Get the underlying openpyxl Workbook object
        """
        return self._workbook
